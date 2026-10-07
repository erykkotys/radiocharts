"""Build the sanitized timetable seed used by the application.

Run from the repository root:
    python scripts/build_timetable_seed.py input.xlsx seed/ramowka_jesien_2026.json

The generated JSON contains no workbook credentials or file-system paths.
"""

from __future__ import annotations

import datetime as dt
import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


WEEKDAY_COLUMNS = {3: 0, 4: 1, 5: 2, 6: 3, 7: 4, 8: 5, 9: 6}
TIME_RE = re.compile(r"(?<!\d)([01]?\d|2[0-3]):([0-5]\d)")

SHOW_NAMES = [
    "5 minut z Bogiem", "Akademia AI", "Bilans – magazyn ekonomiczny",
    "Chartowo na sportowo", "Credo", "Dojrzali na fali", "Fyrtel Sztuki",
    "Gładkie Gotowanie", "Kościół 3.0", "Krag biblijny", "Księga",
    "Lektura w Radiu Emaus", "Lista pokornych", "Lista z moca", "Męskim okiem",
    "Mężczyzną i niewiastą stworzył ich", "Między nami mówiąc",
    "Misyjny Atlas Świata", "Muzyka z celuloidu", "Myśli Świętych",
    "Największy Skarb", "Nie jesteś sam", "Pomost", "Popołudnie z Przewodnikiem",
    "Przedsiebiorcy z Talentem", "Ptaki warte poznania", "Różaniec w Radiu Emaus",
    "Slowo o slowie", "Suma tygodnia", "Święci z nieba ściągnięci",
    "Święta Podróż z Ruchem Focolari", "Trochę Kultury", "U proboszcza w parafii",
    "Wiadomosci z regionu", "Wiadomości Radia Watykańskiego",
    "Wielkopolska nie tylko na weekend", "Wyrwani z niewoli", "Wywiad z czlowiekiem",
    "W drodze do Emaus", "Zapisane w dźwiękach", "Zapytaj księdza",
    "Ze starej plyty", "Zycie z misja", "Zycie na wyspie",
]

ALIASES = {
    "serwis radia watykanskiego": "Wiadomości Radia Watykańskiego",
    "aktualnosci radia watykanskiego": "Wiadomości Radia Watykańskiego",
    "popoludnie z przewodnikiem katolickim": "Popołudnie z Przewodnikiem",
    "przewodnik katolicki": "Popołudnie z Przewodnikiem",
    "focolare": "Święta Podróż z Ruchem Focolari",
    "5 mniut z bogiem": "5 minut z Bogiem",
    "lista z moca": "Lista z moca",
}


def normalized(value: str) -> str:
    source = value.translate(str.maketrans({"ł": "l", "Ł": "L"}))
    text = unicodedata.normalize("NFKD", source).encode("ascii", "ignore").decode().casefold()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


SHOW_LOOKUP = {normalized(name): name for name in SHOW_NAMES}
SHOW_LOOKUP.update(ALIASES)


def clean_title(raw: str) -> str:
    first_line = raw.splitlines()[0].strip()
    first_line = TIME_RE.sub("", first_line, count=1)
    first_line = re.sub(r"^(?:ok(?:oło)?|okolo)\s*", "", first_line, flags=re.I)
    first_line = re.sub(r"\s+", " ", first_line).strip(" -")
    return first_line


def find_shows(title: str) -> list[tuple[str, str]]:
    result: list[tuple[str, str]] = []
    seen: set[str] = set()
    for segment in re.split(r"/", title):
        segment_norm = normalized(segment)
        matches: list[tuple[int, str]] = []
        for alias, canonical in SHOW_LOOKUP.items():
            if re.search(rf"(?:^| ){re.escape(alias)}(?: |$)", segment_norm):
                matches.append((len(alias), canonical))
        for _, canonical in sorted(matches, reverse=True):
            if canonical in seen:
                continue
            seen.add(canonical)
            result.append(
                (canonical, "repeat" if re.search(r"powt", segment, re.I) else "main")
            )
    return result


def classify(title: str) -> tuple[str, str, str]:
    key = normalized(title)
    if "reklam" in key:
        return "Reklamy", "ads", ""
    if "wejsc" in key:
        note = "konkursowe" if "konkurs" in key else ""
        if "zapowiedz mszy" in key:
            note = "zapowiedź mszy"
        return "Wejście prezenterskie", "presenter", note
    if "pogoda" in key or "powitanie" in key:
        return "Pogoda i powitanie", "weather", ""
    if "transmisja mszy" in key:
        return "Transmisja Mszy Świętej", "transmission", ""
    if any(token in key for token in ("apel jasnogorski", "aniol pans", "koronka", "litania")):
        name = re.split(r"\s+dyna|\s+dynamix", title, maxsplit=1, flags=re.I)[0].strip()
        return name, "branding", ""
    if any(token in key for token in ("informacje radia emaus", "serwis informacyjny iar", "podsumowanie dnia")):
        name = re.sub(r"\s*\(.*$", "", title).strip()
        return name, "news", ""
    if "serwis kulturalny" in key or "przeglad prasy" in key or "co przyniesie dzien" in key:
        name = re.sub(r"\s*(?:-|czas).*?$", "", title, flags=re.I).strip()
        return name, "news", ""
    if "muzyczna podroz w czasie" in key:
        return "Muzyczna podróż w czasie – zapowiedź", "branding", ""
    return re.sub(r"\s*\([^)]*\)\s*", " ", title).strip(), "other", ""


def duration_for(title: str, category: str, show_name: str | None = None) -> int:
    minute = re.search(r"(\d+)\s*(?:min(?:uta|uty)?|')", title, re.I)
    if minute:
        value = int(minute.group(1))
        if "2'30" in title:
            return 3
        return max(1, min(value, 240))
    if re.search(r"\d+\s*sek", title, re.I):
        return 1
    if show_name:
        short = normalized(show_name)
        if any(value in short for value in ("5 minut", "mysli swietych", "slowo o slowie")):
            return 5
        if "swieci z nieba" in short or "w drodze do emaus" in short:
            return 10
        return 45
    return {
        "ads": 3,
        "presenter": 1,
        "weather": 2,
        "news": 10,
        "branding": 10,
        "transmission": 60,
        "other": 10,
    }.get(category, 5)


def main(input_path: Path, output_path: Path) -> None:
    workbook = load_workbook(input_path, data_only=True)
    sheet = workbook["Tydzień (kolor)"]
    entries: list[dict[str, object]] = []
    seen: set[tuple[object, ...]] = set()
    last_time = {weekday: "00:00" for weekday in range(7)}

    for row in range(2, sheet.max_row + 1):
        for column, weekday in WEEKDAY_COLUMNS.items():
            raw = sheet.cell(row, column).value
            if not isinstance(raw, str) or not raw.strip():
                continue
            title = clean_title(raw)
            match = TIME_RE.search(raw)
            if match:
                start_time = f"{int(match.group(1)):02d}:{match.group(2)}"
                last_time[weekday] = start_time
            elif normalized(title).startswith("po serwisie"):
                hour = int(last_time[weekday][:2])
                start_time = f"{hour:02d}:03"
            else:
                continue
            approximate = bool(re.search(r"\b(?:około|okolo|ok)\b", raw, re.I))
            shows = find_shows(title)
            candidates: list[dict[str, object]] = []
            if shows:
                for show_name, role in shows:
                    candidates.append(
                        {
                            "name": show_name,
                            "show_name": show_name,
                            "category": "shows",
                            "show_role": role,
                            "note": "",
                            "duration_minutes": duration_for(title, "shows", show_name),
                        }
                    )
            else:
                name, category, note = classify(title)
                if not name:
                    continue
                candidates.append(
                    {
                        "name": name,
                        "category": category,
                        "show_role": "",
                        "note": note,
                        "duration_minutes": duration_for(title, category),
                    }
                )
            for candidate in candidates:
                entry = {
                    **candidate,
                    "weekday": weekday,
                    "start_time": start_time,
                    "approximate": approximate,
                }
                key = (
                    weekday, start_time, candidate["name"], candidate["category"],
                    candidate.get("show_role", ""), candidate.get("note", ""),
                )
                if key in seen:
                    continue
                seen.add(key)
                entries.append(entry)

    payload = {
        "name": "Ramówka Jesień 2026",
        "source_sheet": "Tydzień (kolor)",
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(),
        "entries": entries,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    linked = sum(1 for entry in entries if entry.get("show_name"))
    print(f"Wygenerowano {len(entries)} wpisów, w tym {linked} powiązanych z audycjami.")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        raise SystemExit("Użycie: build_timetable_seed.py INPUT.xlsx OUTPUT.json")
    main(Path(sys.argv[1]), Path(sys.argv[2]))
